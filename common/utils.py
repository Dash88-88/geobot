import os
import random
import re
import shutil
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common.constants import BASE_58_SYMBOLS, RequestReportTitles, RequestReportTotalTitles


def get_base_58_string(length=20):
    return ''.join(random.choices(BASE_58_SYMBOLS, k=length))


def get_current_datetime():
    return datetime.now()


IMAGE_EXTENSIONS = ('.jpg', '.jpeg', '.png', '.webp', '.gif', '.svg', '.bmp', '.ico', '.heic', '.avif', '.tiff')


def is_pure_image_url(text: str) -> bool:
    if not text:
        return False
    clean = text.strip()
    # If text has spaces or multiple lines, it is regular text content
    if ' ' in clean or '\n' in clean or '\t' in clean:
        return False
    if not clean.lower().startswith(('http://', 'https://')):
        return False
    try:
        parsed = urlparse(clean)
        path = parsed.path.lower()
        return path.endswith(IMAGE_EXTENSIONS)
    except Exception:
        return False


def is_url(text: str) -> bool:
    if not text:
        return False
    clean = text.strip()
    return bool(re.match(r'^https?://[^\s/$.?#].[^\s]*$', clean, re.IGNORECASE))



class ReportGenerator:
    root_pwd = Path(os.getcwd(), 'root')
    report_filename = f"BonusRequests-{get_current_datetime().year}.xlsx"
    report_totals_tab_name = 'TOTALS'
    report_filepath = Path(root_pwd, report_filename)

    bold_font = Font(bold=True)
    gray_fill = PatternFill(
        fill_type="solid",
        start_color="DDDDDD",
        end_color="DDDDDD"
    )

    def is_running(self) -> bool:
        file_exists = self.root_pwd.exists()
        if not file_exists:
            os.mkdir(self.root_pwd)
        return file_exists

    def finish(self):
        shutil.rmtree(self.root_pwd)

    def run_bonus_request_generation(self, bonus_requests_data: [], totals_user_data: [], top_referral_sources_data: {}):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove the default sheet

        # Organize data by month
        from collections import defaultdict

        data_by_month = defaultdict(list)
        current_year = get_current_datetime().year

        for record in bonus_requests_data:
            created_at = record.get(RequestReportTitles.request_created_at.value)
            if not isinstance(created_at, datetime):
                continue  # skip invalid dates
            if created_at.year != current_year:
                continue  # skip data not from current year
            month_key = created_at.strftime("%B")
            data_by_month[month_key].append(record)

        headers = [title.value for title in RequestReportTitles]

        for month, records in sorted(data_by_month.items()):
            ws = wb.create_sheet(title=month)
            ws.append(headers)
            ws.freeze_panes = "A2"

            # Header styling
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = self.bold_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.column_dimensions[get_column_letter(col_idx)].width = 22
                cell.fill = self.gray_fill

            # Add records
            for record in records:
                row = [record.get(h, '') for h in headers]
                ws.append(row)

            # Center-align all data cells
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            # Add autofilter
            ws.auto_filter.ref = ws.dimensions
        # Start TOTALS tab creation
        wb.save(self.report_filepath)
        time.sleep(1)

        headers = [title.value for title in RequestReportTotalTitles]
        column_headers = headers[:-2]
        ref_sources_headers = headers[-2:]
        totals_ws = wb.create_sheet(self.report_totals_tab_name)
        for h in column_headers:
            totals_ws.append([h, totals_user_data.get(h, "")])

        totals_ws.append([''])
        totals_ws.append([h for h in ref_sources_headers])

        for row in top_referral_sources_data:
            totals_ws.append([row.get(RequestReportTotalTitles.top_referral_sources.value, ''),
                              row.get(RequestReportTotalTitles.referrals_count.value, '')])

        totals_ws.column_dimensions[get_column_letter(1)].width = 40
        totals_ws.column_dimensions[get_column_letter(2)].width = 22
        totals_ws.column_dimensions[get_column_letter(3)].width = 22

        for row_idx in range(1, len(column_headers)+1):
            cell = totals_ws.cell(row=row_idx, column=1)
            cell.font = self.bold_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

            cell.fill = self.gray_fill

        ref_sources_headers_row = len(column_headers)+2

        for col_idx in range(1, len(ref_sources_headers)+1):
            cell = totals_ws.cell(row=ref_sources_headers_row, column=col_idx)
            cell.font = self.bold_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

            cell.fill = self.gray_fill

        wb.save(self.report_filepath)
        time.sleep(1)
